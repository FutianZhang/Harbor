from pathlib import Path
import inspect
import rewardkit as rk
from rewardkit import criterion as _rk_criterion

_sig = inspect.signature(_rk_criterion)
_supports_id = "id" in _sig.parameters or any(
    p.kind is inspect.Parameter.VAR_KEYWORD for p in _sig.parameters.values()
)

if _supports_id:
    criterion = _rk_criterion
else:
    def criterion(*, id, description, weight=1.0):
        """Adapter for harbor-rewardkit 0.1.7 (no id= / weight= on the decorator)."""
        def deco(fn):
            _rk_criterion(description=description, shared=True)(fn)
            from rewardkit.session import _factory_registry
            _factory_registry[fn.__name__](weight=float(weight), name=str(id))
            return fn
        return deco

import csv
import re
from pathlib import Path

from openpyxl import load_workbook


def norm(value: object) -> str:
    return re.sub(r"\s+", "", "" if value is None else str(value))


def pdf_text(path: Path) -> str:
    try:
        from pypdf import PdfReader
        return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    except Exception:
        return ""


def structured_rows(path: Path, spec: dict) -> tuple[list[str], list[dict[str, str]]]:
    official = {norm(x) for x in spec["ids"]}
    id_header = spec["id_header"]
    if path.suffix.lower() == ".csv":
        try:
            with path.open(encoding="utf-8-sig", newline="") as stream:
                reader = csv.DictReader(stream)
                headers = list(reader.fieldnames or [])
                rows = []
                for row in reader:
                    rid = norm(row.get(id_header) or row.get("finding_id"))
                    if rid in official or any(norm(v) in official for v in row.values()):
                        rows.append(dict(row))
                return headers, rows
        except Exception:
            return [], []
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        rows: list[dict[str, str]] = []
        headers: list[str] = []
        header_needles = {
            norm(id_header),
            "发现id", "治理id", "差距id", "路径id", "findingid", "finding_id",
        }
        for worksheet in workbook.worksheets:
            values = [[cell for cell in row] for row in worksheet.iter_rows(values_only=True)]
            for i, row in enumerate(values):
                current = [norm(v) for v in row]
                if any(x in header_needles for x in current):
                    headers = [str(v or "").strip() for v in row]
                    for data_row in values[i + 1:]:
                        if not any(v is not None and str(v).strip() for v in data_row):
                            continue
                        rec = {
                            headers[j]: "" if j >= len(data_row) or data_row[j] is None else str(data_row[j]).strip()
                            for j in range(len(headers))
                        }
                        blob = {norm(v) for v in rec.values()}
                        if blob & official:
                            rows.append(rec)
                    break
        workbook.close()
        return headers, rows
    except Exception:
        return [], []


def status_column(headers: list[str]) -> str | None:
    for candidate in [
        "风险/状态", "优先级/风险", "状态/风险", "level_status", "status",
        "状态", "风险状态",
    ]:
        if candidate in headers:
            return candidate
    return None


def workbook_sheet_names(path: Path) -> list[str]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        names = list(workbook.sheetnames)
        workbook.close()
        return names
    except Exception:
        return []


def relation_rows(workspace: Path, spec: dict) -> bool:
    structured = workspace / "output" / spec["structured"]
    required_sheets = spec.get("required_sheets", [])
    if structured.suffix.lower() == ".xlsx" and not set(required_sheets).issubset(set(workbook_sheet_names(structured))):
        return False
    headers, rows = structured_rows(structured, spec)
    id_header = spec["id_header"]
    status_header = "优先级/风险" if id_header == "治理ID" else status_column(headers)
    if id_header not in headers or status_header is None:
        return False
    by_id: dict[str, dict[str, str]] = {}
    for row in rows:
        fid = norm(row.get(id_header))
        if not fid or not norm(row.get(status_header)):
            continue
        by_id[fid] = row
    return all(norm(by_id.get(fid, {}).get(status_header)) == norm(status) for fid, status in spec["ids"].items())


def structured_fields(workspace: Path, spec: dict) -> bool:
    structured = workspace / "output" / spec["structured"]
    headers, rows = structured_rows(structured, spec)
    if structured.suffix.lower() == ".csv":
        required = set(spec.get("csv_required_fields", ["evidence", "verification", "rollback"]))
        if not required.issubset(set(headers)):
            return False
        by_id = {norm(row.get(spec["id_header"])): row for row in rows if row.get(spec["id_header"])}
        for fid in spec["ids"]:
            row = by_id.get(fid)
            if not row or any(not norm(row.get(field)) for field in required):
                return False
        return True
    try:
        workbook = load_workbook(structured, read_only=True, data_only=True)
        all_headers: list[set[str]] = []
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows(values_only=True):
                values = {norm(v) for v in row if v is not None and str(v).strip()}
                if values:
                    all_headers.append(values)
        workbook.close()
        if not any("证据" in hs or "evidence" in hs for hs in all_headers):
            return False
        plan = {"责任主体", "完成时限", "验证方式", "回滚方案"}
        plan_en = {"owner", "verification", "rollback"}
        if not any(plan.issubset(hs) or plan_en.issubset(hs) for hs in all_headers):
            return False
    except Exception:
        return False
    by_id = {norm(row.get(spec["id_header"])): row for row in rows if row.get(spec["id_header"])}
    return all(fid in by_id for fid in spec["ids"])


def cross_file_alignment(workspace: Path, spec: dict) -> bool:
    pdf = pdf_text(workspace / "output" / spec["pdf"])
    if not pdf:
        return False
    pdf_norm = norm(pdf)
    structured = workspace / "output" / spec["structured"]
    headers, rows = structured_rows(structured, spec)
    status_header = "优先级/风险" if spec["id_header"] == "治理ID" else status_column(headers)
    if not status_header:
        return False
    by_id: dict[str, dict[str, str]] = {}
    for row in rows:
        fid = norm(row.get(spec["id_header"]))
        if not fid or not norm(row.get(status_header)):
            continue
        by_id[fid] = row
    for fid, status in spec["ids"].items():
        row = by_id.get(fid)
        if not row or norm(row.get(status_header)) != norm(status):
            return False
        pos = pdf_norm.find(norm(fid))
        if pos < 0 or norm(status) not in pdf_norm[pos:pos + 1800]:
            return False
    return True


def derived_constraint(workspace: Path, spec: dict) -> bool:
    if spec.get("derived") == "priority_sum":
        headers, rows = structured_rows(workspace / "output" / spec["structured"], spec)
        if not all(x in headers for x in ["治理ID", "综合分"]):
            return False
        expected = spec["derived_expected"]
        got = {norm(r.get("治理ID")): norm(r.get("综合分")) for r in rows}
        return all(got.get(k) == norm(v) for k, v in expected.items())
    headers, rows = structured_rows(workspace / "output" / spec["structured"], spec)
    ids = [norm(r.get(spec["id_header"])) for r in rows if r.get(spec["id_header"])]
    return bool(ids) and len(ids) == len(set(ids)) and set(ids) == {norm(x) for x in spec["ids"]}



def _read_pdf_text(path: Path) -> str:
    try:
        from pypdf import PdfReader
        return "\n".join((p.extract_text() or "") for p in PdfReader(str(path)).pages)
    except Exception:
        return ""


def _pdf_pages(path: Path) -> int:
    try:
        from pypdf import PdfReader
        return len(PdfReader(str(path)).pages)
    except Exception:
        return 0


def _collect_text(workspace: Path, pdf_name: str) -> str:
    blobs = []
    pdf_path = workspace / "output" / pdf_name
    if pdf_path.is_file():
        blobs.append(_read_pdf_text(pdf_path))
    out = workspace / "output"
    if out.is_dir():
        for p in out.glob("*"):
            if p.suffix.lower() in {".csv", ".xlsx", ".txt", ".md"}:
                try:
                    if p.suffix.lower() == ".xlsx":
                        import openpyxl
                        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
                        for ws in wb.worksheets:
                            blobs.append(ws.title)
                            for row in ws.iter_rows(values_only=True):
                                blobs.append(" ".join("" if v is None else str(v) for v in row))
                        wb.close()
                    else:
                        blobs.append(p.read_text(encoding="utf-8", errors="ignore"))
                except Exception:
                    pass
    return "\n".join(blobs)


def _has_all(text: str, items) -> bool:
    return all(i in text for i in items)


def _id_status_near(text: str, fid: str, status: str, window: int = 1200) -> bool:
    """True only if status appears near this ID, not merely somewhere in the pack."""
    if not text or fid not in text or status not in text:
        return False
    start = 0
    while True:
        i = text.find(fid, start)
        if i < 0:
            break
        lo = max(0, i - window)
        hi = min(len(text), i + len(fid) + window)
        if status in text[lo:hi]:
            return True
        start = i + len(fid)
    return False


def _spec(_workspace=None):
    return SPEC

SPEC = {'pdf': 'CYB-T8-014_ISO27001与等保差距评估报告.pdf', 'structured': 'CYB-T8-014_合规检查与证据矩阵.xlsx', 'id_header': '差距ID', 'ids': {'CMP-001': '不符合/高', 'CMP-002': '不符合/中', 'CMP-003': '证据不足/待复核', 'CMP-004': '改进项/不计入基准缺口', 'CMP-005': '不适用/有范围理由'}, 'required_sheets': ['管理摘要', '控制检查清单', '证据矩阵', '整改路线图']}
PDF_NAME = 'CYB-T8-014_ISO27001与等保差距评估报告.pdf'
STRUCTURED_NAME = 'CYB-T8-014_合规检查与证据矩阵.xlsx'

@criterion(id="P_FILE_PDF", description="主报告 PDF 文件存在且非空", weight=3.0)
def file_pdf(workspace: Path) -> bool:
    p = workspace / "output" / 'CYB-T8-014_ISO27001与等保差距评估报告.pdf'
    return p.is_file() and p.stat().st_size > 100

@criterion(id="P_FILE_XLSX", description="结构化 XLSX 文件存在且非空", weight=3.0)
def file_xlsx(workspace: Path) -> bool:
    p = workspace / "output" / 'CYB-T8-014_合规检查与证据矩阵.xlsx'
    return p.is_file() and p.stat().st_size > 100

@criterion(id="P_XLSX_PARSE", description="结构化交付物为可解析的真实 XLSX（非 Markdown/伪表格）", weight=3.0)
def xlsx_parseable(workspace: Path) -> bool:
    path = workspace / "output" / 'CYB-T8-014_合规检查与证据矩阵.xlsx'
    if not path.is_file() or path.stat().st_size < 32:
        return False
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ok = len(wb.sheetnames) >= 1
        wb.close()
        return ok
    except Exception:
        return False

@criterion(id="P_PDF_PAGES", description="主报告 PDF 页数不少于 9 且可提取文本长度>400", weight=5.0)
def pdf_pages(workspace: Path) -> bool:
    path = workspace / "output" / 'CYB-T8-014_ISO27001与等保差距评估报告.pdf'
    return path.is_file() and _pdf_pages(path) >= 9 and len(_read_pdf_text(path)) > 400

@criterion(id="P_PDF_SECTIONS", description="PDF 含关键章节锚点：管理层摘要/事实与证据链/验证回滚/证据索引", weight=5.0)
def pdf_sections(workspace: Path) -> bool:
    path = workspace / "output" / 'CYB-T8-014_ISO27001与等保差距评估报告.pdf'
    if not path.is_file():
        return False
    text = _read_pdf_text(path)
    return _has_all(text, ['管理层摘要', '事实与证据链', '验证、回滚与恢复门禁', '证据索引'])

@criterion(id="P_CORE_IDS", description="交付物出现全部核心发现/路径 ID：CMP-001,CMP-002,CMP-003,CMP-004,CMP-005", weight=10.0)
def core_ids(workspace: Path) -> bool:
    text = _collect_text(workspace, 'CYB-T8-014_ISO27001与等保差距评估报告.pdf')
    return _has_all(text, ['CMP-001', 'CMP-002', 'CMP-003', 'CMP-004', 'CMP-005'])

@criterion(id="P_STATUS_ANCHORS", description="交付物出现关键状态/结论锚点：不符合/高,不符合/中,证据不足/待复核,改进项/不计入基准缺口,不适用/有范围理由,Not Scored", weight=10.0)
def status_anchors(workspace: Path) -> bool:
    text = _collect_text(workspace, 'CYB-T8-014_ISO27001与等保差距评估报告.pdf')
    return _has_all(text, ['不符合/高', '不符合/中', '证据不足/待复核', '改进项/不计入基准缺口', '不适用/有范围理由', 'Not Scored'])

@criterion(id="P_ID_STATUS_PAIRS", description="每个核心 ID 与其规定状态须在同一发现窗口内同时出现（禁止全文散落或自造编号）", weight=20.0)
def id_status_pairs(workspace: Path) -> bool:
    text = _collect_text(workspace, 'CYB-T8-014_ISO27001与等保差距评估报告.pdf')
    pairs = [('CMP-001', '不符合/高'), ('CMP-002', '不符合/中'), ('CMP-003', '证据不足/待复核'), ('CMP-004', '改进项/不计入基准缺口'), ('CMP-005', '不适用/有范围理由')]
    return all(_id_status_near(text, fid, st) for fid, st in pairs)

@criterion(id="P_DECISION", description="交付物出现任务核心决策锚点：不能宣称全面符合,证据不足,Not Scored", weight=10.0)
def decision_ok(workspace: Path) -> bool:
    text = _collect_text(workspace, 'CYB-T8-014_ISO27001与等保差距评估报告.pdf')
    return _has_all(text, ['不能宣称全面符合', '证据不足', 'Not Scored'])

@criterion(id="P_EVIDENCE_KEYS", description="交付物出现任务特有证据/责任锚点：example_user,TITLE_ID,ISO-A.5.16,CMP-E01,IAM平台主管", weight=10.0)
def evidence_keys(workspace: Path) -> bool:
    text = _collect_text(workspace, 'CYB-T8-014_ISO27001与等保差距评估报告.pdf')
    return _has_all(text, ['example_user', 'TITLE_ID', 'ISO-A.5.16', 'CMP-E01', 'IAM平台主管'])

@criterion(id="P_SHEETS", description="XLSX 工作表名齐全：管理摘要、控制检查清单、证据矩阵、整改路线图", weight=10.0)
def sheets_ok(workspace: Path) -> bool:
    path = workspace / "output" / 'CYB-T8-014_合规检查与证据矩阵.xlsx'
    if not path.is_file():
        return False
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        names = set(wb.sheetnames)
        wb.close()
        return all(s in names for s in ['管理摘要', '控制检查清单', '证据矩阵', '整改路线图'])
    except Exception:
        return False

@criterion(id="P_PAIR_CMP_001", description="官方 ID CMP-001 与规定状态「不符合/高」必须在同一发现窗口内同时出现，禁止全文散落凑齐", weight=15.0)
def pair_cmp_001(workspace: Path) -> bool:
    return _id_status_near(_collect_text(workspace, 'CYB-T8-014_ISO27001与等保差距评估报告.pdf'), 'CMP-001', '不符合/高')

@criterion(id="P_PAIR_CMP_002", description="官方 ID CMP-002 与规定状态「不符合/中」必须在同一发现窗口内同时出现，禁止全文散落凑齐", weight=15.0)
def pair_cmp_002(workspace: Path) -> bool:
    return _id_status_near(_collect_text(workspace, 'CYB-T8-014_ISO27001与等保差距评估报告.pdf'), 'CMP-002', '不符合/中')

@criterion(id="P_PAIR_CMP_003", description="官方 ID CMP-003 与规定状态「证据不足/待复核」必须在同一发现窗口内同时出现，禁止全文散落凑齐", weight=15.0)
def pair_cmp_003(workspace: Path) -> bool:
    return _id_status_near(_collect_text(workspace, 'CYB-T8-014_ISO27001与等保差距评估报告.pdf'), 'CMP-003', '证据不足/待复核')

@criterion(id="P_PAIR_CMP_004", description="官方 ID CMP-004 与规定状态「改进项/不计入基准缺口」必须在同一发现窗口内同时出现，禁止全文散落凑齐", weight=15.0)
def pair_cmp_004(workspace: Path) -> bool:
    return _id_status_near(_collect_text(workspace, 'CYB-T8-014_ISO27001与等保差距评估报告.pdf'), 'CMP-004', '改进项/不计入基准缺口')

@criterion(id="P_PAIR_CMP_005", description="官方 ID CMP-005 与规定状态「不适用/有范围理由」必须在同一发现窗口内同时出现，禁止全文散落凑齐", weight=15.0)
def pair_cmp_005(workspace: Path) -> bool:
    return _id_status_near(_collect_text(workspace, 'CYB-T8-014_ISO27001与等保差距评估报告.pdf'), 'CMP-005', '不适用/有范围理由')


@criterion(id="P_RELATION_ROWS", description="结构化交付物逐行保留全部官方 ID 与规定状态", weight=20.0)
def relation_rows_ok(workspace: Path) -> bool:
    return relation_rows(workspace, _spec(workspace))

@criterion(id="P_GOVERNANCE_FIELDS", description="每个核心 ID 的证据/验证/回滚字段可从结构化清单核验", weight=10.0)
def governance_fields_ok(workspace: Path) -> bool:
    return structured_fields(workspace, _spec(workspace))

@criterion(id="P_CROSS_FILE_ALIGNMENT", description="PDF 与结构化交付物中的官方 ID—状态映射一致", weight=20.0)
def cross_file_ok(workspace: Path) -> bool:
    return cross_file_alignment(workspace, _spec(workspace))

@criterion(id="P_DERIVED_CONSTRAINT", description="任务特有完整 ID 集合满足 Golden 定义", weight=10.0)
def derived_ok(workspace: Path) -> bool:
    return derived_constraint(workspace, _spec(workspace))


@criterion(id="P_COMPLIANCE_IDS", description="合规核验：交付物非空且出现全部官方发现编号（空产物/缺官方 ID 不得在合规维得分）", weight=20.0)
def compliance_ids_ok(workspace: Path) -> bool:
    text = _collect_text(workspace, PDF_NAME)
    if not text or len(text.strip()) < 40:
        return False
    return _has_all(text, list(SPEC["ids"].keys()) if isinstance(SPEC.get("ids"), dict) else list(SPEC["ids"]))

